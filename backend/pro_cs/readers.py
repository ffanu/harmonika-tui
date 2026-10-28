"""Private bounded CPU extraction; extracted document text is untrusted, not visual diagnosis."""
import base64,csv,io,json,re,zipfile,email
from pathlib import Path,PurePosixPath
class ReadError(ValueError):pass
LIMIT=8*1024*1024
EXTENSIONS={'.pdf','.docx','.xlsx','.xls','.pptx','.csv','.json','.xml','.yaml','.yml','.txt','.html','.htm','.md','.png','.jpg','.jpeg','.webp','.tif','.tiff','.svg','.eml','.ics','.vcf','.epub','.zip'}
def archive_check(z):
 infos=z.infolist()
 if len(infos)>2000 or sum(x.file_size for x in infos)>16*1024*1024:raise ReadError('archive_expansion_limit')
 for item in infos:
  p=PurePosixPath(item.filename.replace('\\','/'))
  if p.is_absolute() or '..' in p.parts or ':' in item.filename or item.flag_bits&1 or ((item.external_attr>>16)&0o170000)==0o120000:raise ReadError('unsafe_archive')
  if item.file_size>1024*1024 and item.file_size>max(item.compress_size,1)*100:raise ReadError('archive_compression_limit')

def read(value):
 if type(value) is not dict or set(value)-{'name','data_base64'}:raise ReadError('invalid_upload')
 name=value.get('name');encoded=value.get('data_base64')
 if not isinstance(name,str) or len(name)>160 or Path(name).name!=name or '/' in name or '\\' in name:raise ReadError('invalid_filename')
 if not isinstance(encoded,str) or len(encoded)>12*1024*1024:raise ReadError('upload_limit')
 try:raw=base64.b64decode(encoded,validate=True)
 except Exception:raise ReadError('invalid_base64') from None
 if not 0<len(raw)<=LIMIT:raise ReadError('upload_limit')
 ext=Path(name).suffix.lower()
 if ext not in EXTENSIONS:return {'status':'unavailable','reason':'format_adapter_not_configured','format':ext,'analysis_complete':False}
 mime='application/octet-stream';s='';method='text_extract';structured=None;skipped=0
 if ext in {'.docx','.xlsx','.pptx','.zip','.epub'}:
  with zipfile.ZipFile(io.BytesIO(raw)) as z:archive_check(z)
 if ext=='.pdf':
  if not raw.startswith(b'%PDF-'):raise ReadError('mime_mismatch')
  from pypdf import PdfReader
  doc=PdfReader(io.BytesIO(raw))
  if doc.is_encrypted:raise ReadError('encrypted_document')
  if len(doc.pages)>50:raise ReadError('page_limit')
  s='\n'.join((p.extract_text() or '')[:8000] for p in doc.pages)
  if not s.strip():return {'status':'degraded','reason':'scanned_pdf_requires_ocr_adapter','analysis_complete':False}
 elif ext=='.docx':
  from docx import Document
  doc=Document(io.BytesIO(raw));s='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
 elif ext in ['.xlsx','.xls']:
  cells=[]
  if ext=='.xlsx':
   from openpyxl import load_workbook
   wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
   for sheet in wb:
    for row in sheet.iter_rows(values_only=True):
     if len(cells)>10000:raise ReadError('spreadsheet_limit')
     cells.extend(str(v)[:500] for v in row if v is not None)
   wb.close()
  else:
   import xlrd
   wb=xlrd.open_workbook(file_contents=raw)
   for sheet in wb.sheets():
    if sheet.nrows*sheet.ncols+len(cells)>10000:raise ReadError('spreadsheet_limit')
    cells.extend(str(sheet.cell_value(r,c))[:500] for r in range(sheet.nrows) for c in range(sheet.ncols))
  s='\n'.join(cells)
 elif ext=='.pptx':
  from pptx import Presentation
  deck=Presentation(io.BytesIO(raw))
  if len(deck.slides)>100:raise ReadError('slide_limit')
  s='\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
 elif ext in ['.png','.jpg','.jpeg','.webp','.tif','.tiff']:
  from PIL import Image
  import pytesseract
  Image.MAX_IMAGE_PIXELS=16777216
  with Image.open(io.BytesIO(raw)) as image:
   if image.width*image.height>16777216 or max(image.size)>8192:raise ReadError('pixel_limit')
   if image.format not in ['PNG','JPEG','WEBP','TIFF']:raise ReadError('mime_mismatch')
   image.load();s=pytesseract.image_to_string(image.convert('RGB'),timeout=12)
  method='ocr_only_visual_model_not_invoked'
 elif ext in ['.zip','.epub']:
  from bs4 import BeautifulSoup
  with zipfile.ZipFile(io.BytesIO(raw)) as z:
   items=[]
   for entry in z.infolist():
    if entry.is_dir():continue
    if entry.file_size>1024*1024:raise ReadError('archive_member_limit')
    suffix=PurePosixPath(entry.filename).suffix.lower()
    if suffix in ['.txt','.csv','.md','.html','.xhtml','.json']:
     data=z.read(entry).decode('utf-8',errors='strict');items.append(BeautifulSoup(data,'html.parser').get_text(' ') if suffix in ['.html','.xhtml'] else data)
    else:skipped+=1
    if sum(len(x) for x in items)>32000:raise ReadError('extracted_text_limit')
   s='\n'.join(items)
 else:
  s=raw.decode('utf-8-sig',errors='strict')
  if ext=='.json':structured=json.loads(s);s=json.dumps(structured,ensure_ascii=False)
  elif ext in ['.xml','.svg']:
   from defusedxml.ElementTree import fromstring
   node=fromstring(s);s=' '.join(node.itertext())
  elif ext in ['.yaml','.yml']:
   import yaml
   if s.count('&')+s.count('*')>20:raise ReadError('yaml_alias_limit')
   structured=yaml.safe_load(s);s=json.dumps(structured,ensure_ascii=False)
  elif ext in ['.html','.htm']:
   from bs4 import BeautifulSoup
   doc=BeautifulSoup(s,'html.parser')
   for node in doc(['script','style']):node.decompose()
   s=doc.get_text(' ',strip=True)
  elif ext=='.eml':
   from email import policy
   msg=email.message_from_bytes(raw,policy=policy.default);body=msg.get_body(preferencelist=('plain',));s=str(msg.get('subject',''))+'\n'+(body.get_content() if body else '')
 if len(s.encode())>32000:raise ReadError('extracted_text_limit')
 return {'status':'degraded' if skipped else 'completed','skipped_archive_members':skipped,'format':ext,'text':s,'method':method,'content_is_untrusted':True,'analysis_complete':method!='ocr_only_visual_model_not_invoked' and not skipped,'virus_scan_performed':False,'authenticity_verified':False,'payment_verified':False}
if __name__=='__main__':
 import sys,resource
 resource.setrlimit(resource.RLIMIT_CPU,(20,20));resource.setrlimit(resource.RLIMIT_AS,(2*1024**3,2*1024**3))
 resource.setrlimit(resource.RLIMIT_FSIZE,(32*1024**2,32*1024**2))
 try:
  data=json.loads(sys.stdin.buffer.read(12*1024*1024+1000));out=read(data)
 except Exception as e:out={'status':'degraded','reason':str(e) if isinstance(e,ReadError) else 'file_extraction_failed','analysis_complete':False}
 print(json.dumps(out,ensure_ascii=False))
