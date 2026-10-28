"""Real local fixtures, no customer payloads, no cloud calls or browser launches."""
import argparse,importlib.metadata,json,os,shutil,sqlite3,tempfile,time,subprocess,zipfile
from pathlib import Path
IMPORTS={'fastapi':'fastapi','httpx':'httpx','paramiko':'paramiko','asyncssh':'asyncssh','pypdf':'pypdf','python-docx':'docx','openpyxl':'openpyxl','python-pptx':'pptx','pillow':'PIL','opencv-python-headless':'cv2','faster-whisper':'faster_whisper','redis':'redis','qdrant-client':'qdrant_client','pydantic':'pydantic','cryptography':'cryptography','boto3':'boto3','google-genai':'google.genai'}
def verify():
 checks={};versions={}
 for package,module in IMPORTS.items():
  try:__import__(module);versions[package]=importlib.metadata.version(package);checks[package]=True
  except Exception:checks[package]=False
 with tempfile.TemporaryDirectory(prefix='harmonika-toolchain-fixtures-') as d:
  root=Path(d)
  try:
   from docx import Document
   p=root/'sample.docx';x=Document();x.add_paragraph('Harmonika fixture');x.save(p);checks['docx_roundtrip']=Document(p).paragraphs[0].text=='Harmonika fixture'
   from openpyxl import Workbook,load_workbook
   p=root/'sample.xlsx';x=Workbook();x.active['A1']=42;x.save(p);checks['xlsx_roundtrip']=load_workbook(p).active['A1'].value==42
   from pptx import Presentation
   p=root/'sample.pptx';x=Presentation();x.slides.add_slide(x.slide_layouts[6]);x.save(p);checks['pptx_roundtrip']=len(Presentation(p).slides)==1
   from reportlab.pdfgen.canvas import Canvas
   from pypdf import PdfReader
   p=root/'sample.pdf';x=Canvas(str(p));x.drawString(40,700,'Harmonika fixture');x.save();checks['pdf_roundtrip']='Harmonika fixture' in PdfReader(p).pages[0].extract_text()
   from PIL import Image
   import cv2
   p=root/'sample.png';Image.new('RGB',(96,64),'blue').save(p);checks['image_roundtrip']=cv2.imread(str(p)).shape[:2]==(64,96)
   from PIL import ImageDraw,ImageFont
   import pytesseract
   image=Image.new('RGB',(900,140),'white');font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',64);ImageDraw.Draw(image).text((20,20),'HARMONIKA 123',font=font,fill='black');checks['ocr_real']='HARMONIKA' in pytesseract.image_to_string(image,config='--psm 7',timeout=15).upper()
   video=root/'fixture.mp4';subprocess.run(['ffmpeg','-nostdin','-v','error','-f','lavfi','-i','color=c=blue:s=160x90:r=10:d=0.3','-f','lavfi','-i','sine=frequency=440:duration=0.3','-c:v','mpeg4','-threads','1','-c:a','aac','-shortest',str(video)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=15)
   media=json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-of','json',str(video)],timeout=5));checks['audio_video_real']={x['codec_type'] for x in media['streams']}=={'audio','video'}
   z=root/'fixture.zip'
   with zipfile.ZipFile(z,'w') as f:f.writestr('fixture.txt','Harmonika')
   with zipfile.ZipFile(z) as f:checks['archive_roundtrip']=f.read('fixture.txt')==b'Harmonika'
   c=sqlite3.connect(root/'fixture.sqlite');checks['sqlite_wal']=c.execute('PRAGMA journal_mode=WAL').fetchone()[0]=='wal';c.close()
   from cryptography.fernet import Fernet
   f=Fernet(Fernet.generate_key());checks['crypto_roundtrip']=f.decrypt(f.encrypt(b'fixture'))==b'fixture'
   from qdrant_client import QdrantClient,models
   q=QdrantClient(':memory:');q.create_collection('fixture',vectors_config=models.VectorParams(size=3,distance=models.Distance.COSINE));q.upsert('fixture',[models.PointStruct(id=1,vector=[1.,0.,0.])]);checks['vector_roundtrip']=q.query_points('fixture',query=[1.,0.,0.]).points[0].id==1;q.close()
  except Exception as e:checks['fixture_error']=type(e).__name__
 return {'schema_version':1,'checked_at':int(time.time()),'ok':all(v is True for v in checks.values()),'checks':checks,'versions':versions,'scope':'isolated CPU runtime; no cloud credentials/models/browsers','commands':{x:bool(shutil.which(x)) for x in ['ffmpeg','ffprobe','tesseract','ssh','nmap','docker']},'model_readiness':{'asr':'library_only_model_not_configured','gpu':'external_windows_worker','cloud':'credentials_required'}}
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--output',type=Path);a=p.parse_args();r=verify()
 if a.output:a.output.write_text(json.dumps(r,indent=2));a.output.chmod(0o600)
 print(json.dumps(r));raise SystemExit(0 if r['ok'] else 1)
