"""Local PPTX handler; bounded basic content only."""
from document_common import strings, validate_zip

def inspect(file):
    validate_zip(file)
    from pptx import Presentation
    deck = Presentation(file)
    return {'slides': [{'text': [shape.text for shape in slide.shapes if shape.has_text_frame]} for slide in list(deck.slides)[:100]], 'limits': '100 slides; text frames only'}

def build(data):
    from pptx import Presentation
    result = Presentation()
    slides = data.get('slides')
    if not isinstance(slides, list) or not 1 <= len(slides) <= 100:
        raise ValueError('slides must contain 1..100 slides')
    for spec in slides:
        title = spec.get('title', '')
        if not isinstance(title, str) or len(title) > 100000:
            raise ValueError('title must be a bounded string')
        slide = result.slides.add_slide(result.slide_layouts[1])
        slide.shapes.title.text = title
        body = strings(spec.get('body', []), 'body')
        slide.placeholders[1].text = '\n'.join(body)
    return result
