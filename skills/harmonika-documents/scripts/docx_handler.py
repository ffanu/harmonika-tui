"""Local DOCX handler; bounded basic content only."""
from document_common import strings, validate_zip

def inspect(file):
    validate_zip(file)
    from docx import Document
    doc = Document(file)
    return {'paragraphs': [p.text for p in doc.paragraphs[:1000]], 'tables': [[[c.text for c in row.cells[:50]] for row in table.rows[:100]] for table in doc.tables[:20]], 'limits': '1000 paragraphs, 20 tables, 100 rows/table, 50 columns'}

def build(data):
    from docx import Document
    result = Document()
    for text in strings(data.get('paragraphs'), 'paragraphs'):
        result.add_paragraph(text)
    return result
