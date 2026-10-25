"""Local XLSX handler; bounded basic content only."""
import math
from document_common import strings, validate_zip

def inspect(file):
    validate_zip(file)
    from openpyxl import load_workbook
    book = load_workbook(file, read_only=True, data_only=False, keep_links=False)
    try:
        return {'sheets': [{'name': sheet.title, 'rows': [list(row) for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 0, 200), max_col=min(sheet.max_column or 0, 50), values_only=True)]} for sheet in list(book)[:20]], 'limits': '20 sheets, 200 rows/sheet, 50 columns; formulas returned as text, never evaluated'}
    finally:
        book.close()

def build(data):
    from openpyxl import Workbook
    result = Workbook()
    result.remove(result.active)
    sheets = data.get('sheets')
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= 20:
        raise ValueError('sheets must contain 1..20 sheets')
    names = set()
    for sheet in sheets:
        name = sheet['name']
        if not isinstance(name, str) or not name or len(name) > 31 or any((c in name for c in '[]:*?/\\')) or (name.lower() in names):
            raise ValueError('Invalid or duplicate sheet name')
        names.add(name.lower())
        target = result.create_sheet(name)
        rows = sheet['rows']
        if not isinstance(rows, list) or len(rows) > 1000:
            raise ValueError('rows must be a list of at most 1000 rows')
        for (ri, row) in enumerate(rows, 1):
            if not isinstance(row, list) or len(row) > 100:
                raise ValueError('Each row must have at most 100 cells')
            for (ci, value) in enumerate(row, 1):
                if value is not None and (not isinstance(value, (str, int, float, bool))):
                    raise ValueError('Cell values must be scalar')
                if isinstance(value, float) and (not math.isfinite(value)):
                    raise ValueError('Cell numbers must be finite')
                if isinstance(value, str) and len(value) > 32767:
                    raise ValueError('Cell text too long')
                cell = target.cell(ri, ci, value)
                if isinstance(value, str):
                    cell.data_type = 's'
    return result
