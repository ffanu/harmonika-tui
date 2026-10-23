#!/usr/bin/env python3
"""Minimal local OOXML read/create. No implicit dependency installation."""
import argparse
import importlib.util
import json
import math
import os
from pathlib import Path
import sys
import zipfile

MODULES = {'.docx': 'docx', '.xlsx': 'openpyxl', '.pptx': 'pptx'}
MAX_SIZE = 20 * 1024 * 1024

def capabilities():
    return {ext: bool(importlib.util.find_spec(module)) for ext, module in MODULES.items()}

def scoped(root_value, value):
    root, candidate = Path(root_value), Path(value)
    if not root.is_absolute() or not candidate.is_absolute():
        raise ValueError('Root and file paths must be absolute')
    for p in (root, candidate):
        if '..' in p.parts or any(part.startswith('.') for part in p.parts[1:]):
            raise ValueError('Parent traversal and protected dot-paths are blocked')
        for parent in [p, *p.parents]:
            if parent.is_symlink():
                raise ValueError('Symlink paths are blocked')
    root = root.resolve()
    candidate = candidate.resolve()
    if root in (Path('/'), Path.home()) or not root.is_dir():
        raise ValueError('Root must be an existing project directory')
    if not candidate.is_relative_to(root) or candidate == root:
        raise ValueError('File must be inside the selected project root')
    return candidate

def validate_zip(file):
    if file.stat().st_size > MAX_SIZE:
        raise ValueError('Input exceeds 20 MiB')
    with zipfile.ZipFile(file) as z:
        if len(z.infolist()) > 10000 or sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
            raise ValueError('Expanded document exceeds safe limits')

def dependency(ext):
    if ext not in MODULES:
        raise ValueError('Supported formats: .docx, .xlsx, .pptx')
    if not capabilities()[ext]:
        raise ValueError('Unsupported: missing optional Python module ' + MODULES[ext])

def inspect(file):
    ext = file.suffix.lower()
    dependency(ext)
    validate_zip(file)
    if ext == '.docx':
        from docx import Document
        doc = Document(file)
        return {'paragraphs': [p.text for p in doc.paragraphs[:1000]],
                'tables': [[[c.text for c in row.cells[:50]] for row in table.rows[:100]] for table in doc.tables[:20]],
                'limits': '1000 paragraphs, 20 tables, 100 rows/table, 50 columns'}
    if ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(file, read_only=True, data_only=False, keep_links=False)
        try:
            return {'sheets': [{'name': sheet.title, 'rows': [list(row) for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 0, 200), max_col=min(sheet.max_column or 0, 50), values_only=True)]} for sheet in list(book)[:20]], 'limits': '20 sheets, 200 rows/sheet, 50 columns; formulas returned as text, never evaluated'}
        finally:
            book.close()
    from pptx import Presentation
    deck = Presentation(file)
    return {'slides': [{'text': [shape.text for shape in slide.shapes if shape.has_text_frame]} for slide in list(deck.slides)[:100]], 'limits': '100 slides; text frames only'}

def strings(value, label):
    if not isinstance(value, list) or len(value) > 1000 or any(not isinstance(x, str) or len(x) > 100000 for x in value):
        raise ValueError(label + ' must be a bounded list of strings')
    return value

def create(file, data):
    ext = file.suffix.lower()
    dependency(ext)
    if not isinstance(data, dict):
        raise ValueError('Creation JSON must be an object')
    if ext == '.docx':
        from docx import Document
        result = Document()
        for text in strings(data.get('paragraphs'), 'paragraphs'):
            result.add_paragraph(text)
    elif ext == '.xlsx':
        from openpyxl import Workbook
        result = Workbook()
        result.remove(result.active)
        sheets = data.get('sheets')
        if not isinstance(sheets, list) or not 1 <= len(sheets) <= 20:
            raise ValueError('sheets must contain 1..20 sheets')
        names = set()
        for sheet in sheets:
            name = sheet['name']
            if not isinstance(name, str) or not name or len(name) > 31 or any(c in name for c in '[]:*?/\\') or name.lower() in names:
                raise ValueError('Invalid or duplicate sheet name')
            names.add(name.lower())
            target = result.create_sheet(name)
            rows = sheet['rows']
            if not isinstance(rows, list) or len(rows) > 1000:
                raise ValueError('rows must be a list of at most 1000 rows')
            for ri, row in enumerate(rows, 1):
                if not isinstance(row, list) or len(row) > 100:
                    raise ValueError('Each row must have at most 100 cells')
                for ci, value in enumerate(row, 1):
                    if value is not None and not isinstance(value, (str, int, float, bool)):
                        raise ValueError('Cell values must be scalar')
                    if isinstance(value, float) and not math.isfinite(value):
                        raise ValueError('Cell numbers must be finite')
                    if isinstance(value, str) and len(value) > 32767:
                        raise ValueError('Cell text too long')
                    cell = target.cell(ri, ci, value)
                    if isinstance(value, str):
                        cell.data_type = 's'  # never turn input strings into spreadsheet formulas
    else:
        from pptx import Presentation
        result = Presentation()
        slides = data.get('slides')
        if not isinstance(slides, list) or not 1 <= len(slides) <= 100:
            raise ValueError('slides must contain 1..100 slides')
        for spec in slides:
            title = spec.get('title', '')
            if not isinstance(title, str) or len(title) > 100000:
                raise ValueError('title must be a bounded string')
            slide = result.slides.add_slide(result.slide_layouts[1])
            slide.shapes.title.text = title
            body = strings(spec.get('body', []), 'body')
            slide.placeholders[1].text = '\n'.join(body)
    # Exclusive creation prevents overwriting an existing document.
    if not file.parent.is_dir():
        raise ValueError('Output parent directory does not exist')
    try:
        with file.open('xb') as handle:
            result.save(handle)
    except FileExistsError:
        raise ValueError('Output exists; choose a new filename') from None
    except Exception:
        if file.exists():
            file.unlink()
        raise
    return {'created': str(file), 'bytes': file.stat().st_size, 'format': ext}

def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    sub.add_parser('capabilities')
    for command in ('inspect', 'create'):
        p = sub.add_parser(command)
        p.add_argument('--root', required=True)
        p.add_argument('--input' if command == 'inspect' else '--output', required=True)
        if command == 'create':
            p.add_argument('--data', required=True)
    args = parser.parse_args(argv)
    if args.command == 'capabilities':
        return capabilities()
    if args.command == 'inspect':
        return inspect(scoped(args.root, args.input))
    output = scoped(args.root, args.output)
    source = scoped(args.root, args.data)
    if source.stat().st_size > 1024 * 1024:
        raise ValueError('Creation JSON exceeds 1 MiB')
    data = json.loads(source.read_text(encoding='utf-8'))
    return create(output, data)

if __name__ == '__main__':
    # Optional runtime installed explicitly by the Harmonika owner; never install it here.
    runtime_root = Path.home() / '.config/harmonika-terminal/document-runtime'
    runtime_python = runtime_root / 'bin/python'
    if runtime_python.is_file() and Path(sys.prefix).resolve() != runtime_root.resolve() and not all(capabilities().values()):
        os.execv(str(runtime_python), [str(runtime_python), str(Path(__file__).resolve()), *sys.argv[1:]])
    try:
        print(json.dumps(main(), ensure_ascii=False, default=str, allow_nan=False))
    except Exception as error:
        print(json.dumps({'error': str(error)}, ensure_ascii=False), file=sys.stderr)
        sys.exit(2)
